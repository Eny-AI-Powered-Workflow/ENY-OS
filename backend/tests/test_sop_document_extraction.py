# /home/obed/Documents/Eny_consulting/Eny_consulting/backend/tests/test_sop_document_extraction.py

from io import BytesIO

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.services import sop_document_service
from app.services.sop_document_service import DocumentExtractionError, extract_sop_text


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Lead qualification procedure")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Lead status"
    table.cell(0, 1).text = "Next action"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "Enrollment follow-up"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "Status"
    table.cell(0, 1).text = "Action"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_extracts_word_paragraphs_and_tables():
    extracted = extract_sop_text("qualification.DOCX", _docx_bytes())

    assert "Lead qualification procedure" in extracted
    assert "Lead status | Next action" in extracted


def test_extracts_powerpoint_slide_text_and_tables():
    extracted = extract_sop_text("enrollment.pptx", _pptx_bytes())

    assert "## Slide 1" in extracted
    assert "Enrollment follow-up" in extracted
    assert "Status | Action" in extracted


def test_extracts_pdf_text_with_page_context(monkeypatch):
    class FakePage:
        def extract_text(self):
            return "Escalate urgent leads."

    class FakePdfReader:
        is_encrypted = False
        pages = [FakePage()]

        def __init__(self, *_args, **_kwargs):
            pass

    monkeypatch.setattr(sop_document_service, "PdfReader", FakePdfReader)

    assert extract_sop_text("procedure.pdf", b"%PDF-1.7") == "## Page 1\n\nEscalate urgent leads."


def test_pdf_without_extractable_text_explains_ocr_requirement(monkeypatch):
    class FakePage:
        def extract_text(self):
            return None

    class FakePdfReader:
        is_encrypted = False
        pages = [FakePage()]

        def __init__(self, *_args, **_kwargs):
            pass

    monkeypatch.setattr(sop_document_service, "PdfReader", FakePdfReader)

    with pytest.raises(DocumentExtractionError, match="Run OCR"):
        extract_sop_text("scanned.pdf", b"%PDF-1.7")


def test_rejects_mismatched_file_content():
    with pytest.raises(DocumentExtractionError, match="do not match a Word"):
        extract_sop_text("procedure.docx", b"%PDF-1.7")


def test_rejects_legacy_word_documents_with_conversion_guidance():
    with pytest.raises(DocumentExtractionError, match="Save the file as .docx"):
        extract_sop_text("procedure.doc", b"legacy word")
