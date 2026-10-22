# /home/obed/Documents/Eny_consulting/Eny_consulting/backend/app/services/sop_document_service.py

from io import BytesIO
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml.etree import XMLSyntaxError
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


MAX_OFFICE_PACKAGE_BYTES = 100 * 1024 * 1024


class DocumentExtractionError(ValueError):
    """Raised when a supported SOP file cannot provide readable text."""


def _validate_office_package(content: bytes, required_member: str, label: str) -> None:
    try:
        with ZipFile(BytesIO(content)) as package:
            members = package.infolist()
            if required_member not in {member.filename for member in members}:
                raise DocumentExtractionError(f"This file is not a valid {label} document.")
            if sum(member.file_size for member in members) > MAX_OFFICE_PACKAGE_BYTES:
                raise DocumentExtractionError(
                    f"The uncompressed {label} document exceeds the 100 MB processing limit."
                )
    except BadZipFile as exc:
        raise DocumentExtractionError(f"This file is not a valid {label} document.") from exc


def _docx_text(content: bytes) -> str:
    _validate_office_package(content, "word/document.xml", "Word")
    try:
        document = Document(BytesIO(content))
        sections: list[str] = []
        for block in document.element.body.iterchildren():
            if isinstance(block, CT_P):
                paragraph = Paragraph(block, document).text.strip()
                if paragraph:
                    sections.append(paragraph)
            elif isinstance(block, CT_Tbl):
                table = Table(block, document)
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        sections.append(" | ".join(cells))
    except (DocxPackageNotFoundError, BadZipFile, KeyError, OSError, ValueError, XMLSyntaxError) as exc:
        raise DocumentExtractionError("The Word document could not be read.") from exc
    return "\n\n".join(sections)


def _pptx_text(content: bytes) -> str:
    _validate_office_package(content, "ppt/presentation.xml", "PowerPoint")
    try:
        presentation = Presentation(BytesIO(content))
        slides: list[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_text: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    slide_text.append(shape.text.strip())
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            slide_text.append(" | ".join(cells))
            if slide_text:
                slides.append(f"## Slide {slide_number}\n\n" + "\n\n".join(slide_text))
    except (PptxPackageNotFoundError, BadZipFile, KeyError, OSError, ValueError, XMLSyntaxError) as exc:
        raise DocumentExtractionError("The PowerPoint presentation could not be read.") from exc
    return "\n\n".join(slides)


def _pdf_text(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise DocumentExtractionError(
                "This PDF is password-protected. Remove its password before uploading."
            )
        pages = []
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                pages.append(f"## Page {page_number}\n\n{page_text}")
    except DocumentExtractionError:
        raise
    except (PdfReadError, OSError, ValueError) as exc:
        raise DocumentExtractionError("The PDF could not be read.") from exc
    return "\n\n".join(pages)


def extract_sop_text(filename: str, content: bytes) -> str:
    """Extract readable text from an uploaded PDF, DOCX, or PPTX SOP."""
    if not content:
        raise DocumentExtractionError("The uploaded document is empty.")

    extension = PurePath(filename).suffix.lower()
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise DocumentExtractionError("The file contents do not match a PDF document.")
        extracted = _pdf_text(content)
    elif extension == ".docx":
        if not content.startswith(b"PK"):
            raise DocumentExtractionError("The file contents do not match a Word document.")
        extracted = _docx_text(content)
    elif extension == ".pptx":
        if not content.startswith(b"PK"):
            raise DocumentExtractionError("The file contents do not match a PowerPoint presentation.")
        extracted = _pptx_text(content)
    elif extension == ".doc":
        raise DocumentExtractionError("Legacy Word .doc files are not supported. Save the file as .docx.")
    else:
        raise DocumentExtractionError("Supported document uploads are PDF, Word (.docx), and PowerPoint (.pptx).")

    normalized = extracted.replace("\x00", "").strip()
    if not normalized:
        if extension == ".pdf":
            raise DocumentExtractionError(
                "No selectable text was found in this PDF. Run OCR to make scanned pages searchable, then upload it again."
            )
        raise DocumentExtractionError("No readable text was found in this document.")
    return normalized
